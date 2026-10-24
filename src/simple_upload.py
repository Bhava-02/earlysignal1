"""
The easy way for a college to bring its own data: ONE Excel workbook, the same Roll No in every sheet.

Only three sheets are needed (Students, Marks, Attendance). Semesters, the current week and the list of
departments are worked out from the data, so the college only types its name.

The workbook is checked, then converted into the same internal files the sample college uses, and the
normal pipeline runs on them. Rows whose Roll No starts with EXAMPLE are ignored.
"""
import io
import math
import re

import numpy as np
import pandas as pd

from .integrate import DEPT_ALIASES, norm_roll

LABELS = {-2: "Two semesters ago", -1: "Last semester", 0: "This semester"}
MIN_WEEK_NOW = 6     # enough weeks this semester to see trends
MIN_PAST_WEEKS = 12  # a past semester must be (nearly) complete


def key(s):
    return re.sub(r"[^a-z0-9]", "", str(s).lower())


# column: (needed?, what to put, other names a college might use)
C_ = lambda need, help_, *alias: (need, help_, alias)  # noqa: E731
SHEETS = {
    "Students": dict(
        required=True, per="One row per student", why="Who the students are",
        missing="Nothing works without the list of students.",
        cols={
            "Roll No": C_(True, "The student's roll number. Use the same one in every sheet.",
                          "rollnumber", "roll", "regno", "registerno", "registernumber", "studentid", "usn",
                          "enrollmentno", "enrolmentno", "admissionno", "id"),
            "Name": C_(True, "Student's name", "studentname", "fullname"),
            "Department": C_(True, "CSE, ECE, MECH ... (any names, used as they are)", "dept", "branch", "programme",
                             "program", "course"),
            "Current semester": C_(False, "1 to 8. If blank, we take the latest semester in Attendance.",
                                   "currentsem", "semester", "sem", "semesterno"),
            "Section": C_(False, "A, B, C ...", "sec", "class"),
            "Gender": C_(False, "Only used to check fairness, never to score", "sex"),
            "Hostel or day scholar": C_(False, "Hostel / Day Scholar (fairness check only)", "residence", "hostel",
                                        "accommodation", "hostelday", "dayscholar"),
            "First graduate": C_(False, "Yes / No (fairness check only)", "firstgraduate", "firstgen",
                                 "firstgeneration", "firstgenerationlearner"),
            "Mentor": C_(False, "Mentor's name", "mentorname", "facultyadvisor", "advisor", "counsellor"),
            "Email": C_(False, "College email", "emailid", "mail", "officialemail"),
            "12th mark %": C_(False, "Used only when a student has no CGPA yet", "12th", "12thmark", "12thmarks",
                              "12thpercentage", "hsc", "hscpercent", "hscmark", "class12"),
        }),
    "Marks": dict(
        required=True, per="One row per student, subject and semester", why="CGPA, arrears and internal test marks",
        missing="Needed for CGPA, arrears and internal marks, and to learn who failed in past semesters.",
        cols={
            "Roll No": C_(True, "Same roll number as in Students", "rollnumber", "roll", "regno", "registerno",
                          "registernumber", "studentid", "usn", "id"),
            "Semester": C_(True, "Semester number, 1 to 8", "sem", "semesterno", "semesternumber"),
            "Subject code": C_(True, "e.g. CS301", "subject", "coursecode", "papercode", "subcode", "code"),
            "Subject name": C_(False, "e.g. Data Structures", "coursename", "papername", "subjecttitle", "title"),
            "Internal test 1": C_(True, "Marks (out of 50 or 100, we detect which). AB if absent.", "cia1", "internal1",
                                  "ia1", "test1", "mid1", "midterm1", "cat1", "internal"),
            "Internal test 2": C_(False, "Marks, same scale as test 1", "cia2", "internal2", "ia2", "test2", "mid2",
                                  "midterm2", "cat2"),
            "End semester": C_(False, "End-semester exam marks", "endsem", "external", "semesterexam", "endsemester",
                               "ese", "see", "university"),
            "Grade": C_(True, "Grade points 0-10, or the letter (O, A+, A, B+, B, C, U). Blank for this semester.",
                        "gradepoint", "gradepoints", "gp", "result", "lettergrade"),
            "Attempt": C_(False, "1 = first attempt (default). 2 or more = arrear exam.", "attemptno", "attempts"),
        }),
    "Attendance": dict(
        required=True, per="One row per student per week (per subject if you have it)",
        why="Attendance and the weekly early warning",
        missing="Needed for attendance, shortage warnings and the weekly risk model.",
        cols={
            "Roll No": C_(True, "Same roll number as in Students", "rollnumber", "roll", "regno", "registerno",
                          "registernumber", "studentid", "usn", "id"),
            "Semester": C_(True, "Semester number, 1 to 8", "sem", "semesterno", "semesternumber"),
            "Week": C_(True, "Week of the semester: 1, 2, 3 ...", "weekno", "weeknumber"),
            "Subject code": C_(False, "Optional", "subject", "coursecode", "papercode", "subcode", "code"),
            "Classes held": C_(True, "Number of classes that week", "held", "totalclasses", "conducted",
                               "classesconducted", "periodsconducted", "total"),
            "Classes attended": C_(True, "Number the student attended", "attended", "present", "classespresent",
                                   "periodsattended", "periodspresent"),
        }),
    "LMS": dict(
        required=False, per="One row per student per week", why="Online activity and assignments",
        missing="Engagement and early warnings use attendance and marks only.",
        cols={
            "Roll No": C_(True, "Same roll number as in Students", "rollnumber", "roll", "regno", "studentid", "id"),
            "Semester": C_(True, "Semester number", "sem", "semesterno"),
            "Week": C_(True, "Week of the semester", "weekno", "weeknumber"),
            "Logins": C_(True, "Logins that week", "lmslogins", "loginsthisweek", "logincount"),
            "Minutes online": C_(False, "Minutes spent that week", "minutes", "timespent", "timeonline"),
            "Assignments due": C_(False, "Assignments due that week", "assignmentsgiven", "due"),
            "Assignments submitted": C_(False, "Assignments the student submitted", "submitted", "assignmentsdone"),
            "Quizzes taken": C_(False, "Quiz attempts that week", "quizattempts", "quizzes"),
        }),
    "Other details": dict(
        required=False, per="One row per student. Fill any columns you have; leave the rest blank.",
        why="Placement training, skills, activities and feedback",
        missing="Placement readiness, activities and feedback count as unknown.",
        cols={
            "Roll No": C_(True, "Same roll number as in Students", "rollnumber", "roll", "regno", "studentid", "id"),
            "Aptitude score": C_(False, "Best aptitude test score, out of 100", "aptitude"),
            "Coding score": C_(False, "Best coding test score, out of 100", "coding"),
            "Mock interviews done": C_(False, "Count", "mockinterviews", "mocks", "mockinterviewcount"),
            "Mock interview rating": C_(False, "Average rating, 1 to 5", "mockrating", "mockinterviewscore"),
            "Technical skill score": C_(False, "Out of 100", "technical", "technicalscore", "technicalskills"),
            "Soft skill score": C_(False, "Out of 100", "softskill", "softskills", "softskillsscore", "communication"),
            "Certifications": C_(False, "Count", "certificates", "certificationcount"),
            "Events and clubs": C_(False, "Count in the last 12 months", "events", "clubs", "activities"),
            "Hackathons": C_(False, "Count in the last 12 months", "hackathon"),
            "Faculty rating": C_(False, "Class participation, 1 to 5", "participationrating", "facultyfeedback",
                                 "participation"),
            "Faculty remark": C_(False, "Short comment", "remark", "remarks", "comment"),
            "Course satisfaction": C_(False, "Survey, 1 to 5", "satisfaction"),
            "Feels supported": C_(False, "Survey, 1 to 5", "supported", "support"),
            "Workload stress": C_(False, "Survey, 1 to 5", "stress", "workload"),
        }),
    "Past actions": dict(
        required=False, per="One row per action taken", why="Learning which actions work",
        missing="'What works' starts once actions are logged in the app.",
        cols={
            "Roll No": C_(True, "Same roll number as in Students", "rollnumber", "roll", "regno", "studentid", "id"),
            "Semester": C_(True, "Semester number", "sem", "semesterno"),
            "Week": C_(True, "Week of the semester", "weekno"),
            "Action": C_(True, "e.g. Mentor meeting, Remedial classes, Parent call", "actiontaken", "intervention"),
            "Done by": C_(False, "Who did it", "owner", "by", "staff", "faculty"),
        }),
    "Past placements": dict(
        required=False, per="One row per graduate from last year", why="What gets students placed at your college",
        missing="Placement predictions use the built-in sample history.",
        cols={
            "Student ID": C_(False, "Any ID, optional", "rollno", "alumniid", "id", "regno"),
            "Department": C_(False, "", "dept", "branch"),
            "Final CGPA": C_(True, "", "cgpa"),
            "Standing arrears": C_(False, "Count at graduation", "arrears", "backlogs", "standingbacklogs"),
            "Aptitude score": C_(False, "Out of 100", "aptitude", "bestaptitude"),
            "Coding score": C_(False, "Out of 100", "coding", "bestcoding"),
            "Technical skill score": C_(False, "Out of 100", "technical", "technicalscore"),
            "Soft skill score": C_(False, "Out of 100", "softskill", "softskills"),
            "Mock interviews done": C_(False, "Count", "mockinterviews", "mocks"),
            "Mock interview rating": C_(False, "1 to 5", "mockrating"),
            "Certifications": C_(False, "Count", "certificates"),
            "Internship": C_(False, "Yes / No", "didinternship"),
            "Placed": C_(True, "Yes / No", "gotplaced", "placementstatus", "placedyesno", "status"),
        }),
}
REQUIRED = [s for s, v in SHEETS.items() if v["required"]]

EXAMPLES = {
    "Students": [
        ["EXAMPLE-1", "Asha R", "CSE", 5, "A", "F", "Hostel", "No", "Dr. Meena", "asha@college.edu", 91],
        ["EXAMPLE-2", "Karthik", "ECE", 3, "B", "M", "Day Scholar", "Yes", "Prof. Ravi", "karthik@college.edu", 78],
        ["EXAMPLE-3", "Divya P", "MECH", 7, "A", "F", "Day Scholar", "No", "Dr. Arun", "divya@college.edu", 84]],
    "Marks": [
        ["EXAMPLE-1", 3, "CS301", "Data Structures", 38, 41, 72, "A", 1],
        ["EXAMPLE-1", 4, "CS402", "Operating Systems", 29, 33, 45, "U", 1],
        ["EXAMPLE-1", 4, "CS402", "Operating Systems", None, None, 58, "B", 2],
        ["EXAMPLE-1", 5, "CS501", "Computer Networks", 35, None, None, None, 1]],
    "Attendance": [
        ["EXAMPLE-1", 5, 1, "CS501", 4, 4], ["EXAMPLE-1", 5, 2, "CS501", 4, 3], ["EXAMPLE-1", 4, 1, "CS402", 4, 4]],
    "LMS": [["EXAMPLE-1", 5, 1, 6, 140, 2, 2, 1], ["EXAMPLE-1", 5, 2, 3, 60, 2, 1, 0]],
    "Other details": [
        ["EXAMPLE-1", 72, 65, 2, 3.5, 70, 64, 1, 3, 1, 4, "Asks good questions", 4, 4, 3],
        ["EXAMPLE-2", 58, None, 0, None, 61, 55, 0, 1, 0, 3, None, 3, 2, 4]],
    "Past actions": [["EXAMPLE-1", 4, 7, "Mentor meeting", "Dr. Meena"],
                     ["EXAMPLE-2", 2, 9, "Remedial classes", "Prof. Ravi"]],
    "Past placements": [
        ["EXAMPLE-A1", "CSE", 8.1, 0, 78, 72, 74, 70, 3, 4, 2, "Yes", "Yes"],
        ["EXAMPLE-A2", "MECH", 6.2, 1, 55, 30, 52, 60, 0, None, 0, "No", "No"]],
}


# ---------------------------------------------------------------- reading
def read_files(files):
    """files: list of (file name, bytes). Returns {sheet: DataFrame} and a list of notes about unrecognised input."""
    found, notes = {}, []
    for name, data in files:
        try:
            if name.lower().endswith((".xlsx", ".xlsm", ".xls")):
                tables = pd.read_excel(io.BytesIO(data), sheet_name=None, dtype=object)
            else:
                tables = {re.sub(r"\.[^.]+$", "", name): pd.read_csv(io.BytesIO(data), dtype=object)}
        except Exception as e:  # unreadable file
            notes.append(f"Could not read {name}: {e}")
            continue
        for tab, df in tables.items():
            if key(tab) in ("starthere", "readme", "instructions"):
                continue
            df = df.dropna(how="all")
            sheet = identify(tab, df)
            if sheet is None:
                if len(df):
                    notes.append(f"'{tab}' in {name} was not recognised (its sheet name and columns don't match the "
                                 f"template), so it was skipped.")
                continue
            found[sheet] = pd.concat([found[sheet], df]) if sheet in found else df
    return found, notes


def identify(tab, df):
    """Which template sheet is this? Exact sheet/file name, else its columns, else a partial name."""
    k = key(tab)
    for sheet in SHEETS:
        if k == key(sheet):
            return sheet
    best, best_n = None, 0
    for sheet in SHEETS:  # every needed column present; prefer the sheet that needs the most of them
        need = [c for c, (n, *_r) in SHEETS[sheet]["cols"].items() if n]
        if all(match_column(df.columns, sheet, c) is not None for c in need) and len(need) > best_n:
            best, best_n = sheet, len(need)
    if best:
        return best
    for sheet in SHEETS:
        if key(sheet) in k:
            return sheet
    return None


def match_column(columns, sheet, col):
    _, _, alias = SHEETS[sheet]["cols"][col]
    wanted = [key(col)] + list(alias)
    keys = {key(c): c for c in columns}
    for w in wanted:
        if w in keys:
            return keys[w]
    return None


def standardise(sheet, df):
    """Rename the college's columns to the template's names. Returns (df, missing needed columns)."""
    out, missing, used = pd.DataFrame(index=df.index), [], set()
    for col, (need, _, _) in SHEETS[sheet]["cols"].items():
        src = match_column([c for c in df.columns if c not in used], sheet, col)
        if src is None:
            out[col] = np.nan
            if need:
                missing.append(col)
        else:
            used.add(src)
            out[col] = df[src]
    return out.reset_index(drop=True), missing


# ---------------------------------------------------------------- cleaning helpers
def num(s):
    return pd.to_numeric(pd.Series(s).astype(str).str.replace("%", "", regex=False).str.strip()
                         .replace({"": None, "nan": None, "None": None}), errors="coerce")


def clean_dept(label):
    if pd.isna(label) or not str(label).strip():
        return None
    k = re.sub(r"[^a-z&]", "", str(label).lower()).replace("&", "")
    return DEPT_ALIASES.get(k, re.sub(r"\s+", " ", str(label).strip()).upper())


LETTER = {"O": "O", "S": "O", "A+": "A+", "A": "A", "B+": "B+", "B": "B", "C": "C", "D": "C", "E": "C", "P": "C",
          "PASS": "C", "U": "RA", "RA": "RA", "F": "RA", "FAIL": "RA", "AB": "RA", "ABSENT": "RA", "R": "RA",
          "W": "RA", "I": "RA", "X": "RA", "WH": "RA"}


def grade(v):
    """Grade points (0-10) or a letter grade -> the letter scale the models use."""
    if pd.isna(v) or not str(v).strip():
        return np.nan
    s = str(v).strip().upper()
    try:
        p = float(s)
    except ValueError:
        return LETTER.get(s.replace(" ", ""), "?")
    p = round(p)
    return {10: "O", 9: "A+", 8: "A", 7: "B+", 6: "B", 5: "C", 4: "C"}.get(p, "RA" if 0 <= p < 4 else "?")


def yes_no(v):
    s = str(v).strip().lower()
    if s in ("yes", "y", "1", "1.0", "true", "placed"):
        return "Yes"
    if s in ("no", "n", "0", "0.0", "false", "not placed", "notplaced"):
        return "No"
    return np.nan


def out_of(values):
    """Guess the maximum mark of internal tests from the highest mark anyone got."""
    top = values.max()
    if pd.isna(top):
        return 50
    for d in (10, 20, 25, 30, 40, 50, 60, 75, 80, 100):
        if top <= d:
            return d
    return float(top)


def drop_examples(df, col="Roll No"):
    ex = df[col].astype(str).str.strip().str.upper().str.startswith("EXAMPLE")
    return df[~ex].reset_index(drop=True)


# ---------------------------------------------------------------- check + convert
def prepare(raw):
    """
    raw: {sheet: DataFrame} as read from the upload.
    Returns dict(checks=[...], ok=bool, files={internal file name: DataFrame}, settings={...}, facts={...}).
    Each check is dict(sheet, status in ok / warn / skip / error, text).
    """
    checks, std = [], {}

    def add(sheet, status, text):
        checks.append(dict(sheet=sheet, status=status, text=text))

    for sheet, spec in SHEETS.items():
        if sheet not in raw:
            add(sheet, "error" if spec["required"] else "skip",
                "Missing. This sheet is needed." if spec["required"] else f"Not included. {spec['missing']}")
            continue
        df, missing = standardise(sheet, raw[sheet])
        if missing:
            add(sheet, "error", "Missing column(s): " + ", ".join(missing) + ".")
            continue
        df = drop_examples(df, "Student ID" if sheet == "Past placements" else "Roll No")
        if sheet != "Past placements":
            df["roll"] = df["Roll No"].map(norm_roll)
            df = df[df["roll"].notna()]
        if not len(df):
            add(sheet, "error" if spec["required"] else "skip",
                "Only example rows, no real data yet." if spec["required"] else "Only example rows, so it was skipped.")
            continue
        std[sheet] = df
    result = dict(checks=checks, ok=False, files={}, settings={}, facts={})
    if any(c["status"] == "error" for c in checks):
        return result

    # ---- Students
    S = std["Students"].copy()
    dup = int(S.duplicated("roll").sum())
    S = S.drop_duplicates("roll")
    S["dept"] = S["Department"].map(clean_dept)
    no_dept = int(S["dept"].isna().sum())
    S["dept"] = S["dept"].fillna("UNKNOWN")
    roll_set = set(S["roll"])

    # ---- Marks and Attendance: roll numbers and semester numbers
    def link(sheet):
        d = std[sheet].copy()
        unknown = ~d["roll"].isin(roll_set)
        d["sem"] = num(d["Semester"]).values
        bad = d["sem"].isna() | (d["sem"] < 1) | (d["sem"] > 12)
        return d[~unknown & ~bad].copy(), int(unknown.sum()), int((bad & ~unknown).sum())

    M, m_unknown, m_badsem = link("Marks")
    A, a_unknown, a_badsem = link("Attendance")
    A["week"] = num(A["Week"]).values
    A["held"] = num(A["Classes held"]).values
    A["attd"] = num(A["Classes attended"]).values
    a_badnum = int((A["week"].isna() | A["held"].isna() | A["attd"].isna()).sum())
    A = A[A["week"].notna() & A["held"].notna() & A["attd"].notna()]

    # current semester of each student: as given, else the latest semester seen in Attendance (then Marks)
    S["cur"] = num(S["Current semester"]).values
    given_cur = int(S["cur"].notna().sum())
    S["cur"] = S["cur"].fillna(S["roll"].map(A.groupby("roll")["sem"].max())).fillna(
        S["roll"].map(M.groupby("roll")["sem"].max()))
    no_cur = int(S["cur"].isna().sum())
    S = S[S["cur"].notna()].copy()
    S["year"] = np.clip(np.ceil(S["cur"] / 2), 1, 4).astype(int)
    cur = S.set_index("roll")["cur"]
    A["rel"] = A["sem"] - A["roll"].map(cur)
    M["rel"] = M["sem"] - M["roll"].map(cur)
    A = A[A["rel"].isin([-2, -1, 0])]
    A["semester"] = A["rel"].map(LABELS)

    # ---- the semesters we need
    def weeks(rel):
        g = A[A["rel"] == rel].groupby("roll")["week"].max()
        return int(g.mode().iloc[0]) if len(g) else 0, int(len(g))

    (w_now, n_now), (w_last, n_last), (w_prev, n_prev) = weeks(0), weeks(-1), weeks(-2)
    problems = []
    if not n_now:
        problems.append("Attendance has no rows for the current semester.")
    elif w_now < MIN_WEEK_NOW:
        problems.append(f"This semester's attendance only goes up to week {w_now}. Upload from week {MIN_WEEK_NOW} "
                        f"onward, when there is enough to see trends.")
    for w, n, what in [(w_last, n_last, "last semester"), (w_prev, n_prev, "the semester before that")]:
        if not n:
            problems.append(f"Attendance needs {what} too. The early warning learns from the two past semesters.")
        elif w < MIN_PAST_WEEKS:
            problems.append(f"Attendance for {what} stops at week {w}. Please include the whole semester.")

    M["grade"] = M["Grade"].map(grade)
    bad_grade = int((M["grade"] == "?").sum())
    M.loc[M["grade"] == "?", "grade"] = np.nan
    M["attempt"] = num(M["Attempt"]).fillna(1).astype(int).values
    first = M[(M["attempt"] == 1) & M["grade"].notna()]
    for rel, what in [(-1, "last semester"), (-2, "the semester before that")]:
        g = first[first["rel"] == rel]
        fails = g.groupby("roll")["grade"].apply(lambda x: (x == "RA").any()).sum() if len(g) else 0
        if not len(g):
            problems.append(f"Marks need grades for {what}, so the model can learn who failed.")
        elif fails < 5 or fails > g["roll"].nunique() - 5:
            problems.append(f"Marks for {what} have {fails} student(s) with a failed subject. The model needs at "
                            f"least 5 who failed and 5 who passed to learn from.")
    for p in problems:
        add("Semesters", "error", p)

    # internal test scale
    ia1 = M["Internal test 1"]
    absent = ia1.astype(str).str.strip().str.upper().isin(["AB", "ABSENT", "A"])
    scale = 50 / out_of(num(ia1[~absent]))
    cia1 = (num(ia1) * scale).round(1).astype(object)
    cia1[absent.values] = "AB"
    max_mark = out_of(num(ia1[~absent]))

    # ---- report on the required sheets
    n_depts = S["dept"].nunique()
    yrs = S["year"].value_counts().sort_index()
    add("Students", "ok", f"{len(S):,} students in {n_depts} department(s): "
        + ", ".join(f"{d} ({n})" for d, n in S["dept"].value_counts().items()) + ". By year: "
        + ", ".join(f"year {y}: {n}" for y, n in yrs.items()) + ".")
    if dup:
        add("Students", "warn", f"{dup} duplicate roll number(s); kept the first.")
    if no_dept:
        add("Students", "warn", f"{no_dept} student(s) have no department; shown as UNKNOWN.")
    if no_cur:
        add("Students", "warn", f"{no_cur} student(s) have no current semester and no attendance; left out.")
    if given_cur < len(S) + no_cur:
        add("Students", "ok", "Current semester taken from Attendance where it was blank.")
    add("Marks", "ok", f"{len(M):,} marks. Internal tests look like they are out of {max_mark:g}.")
    if m_unknown:
        add("Marks", "warn", f"{m_unknown} row(s) have a roll number that is not in Students; skipped.")
    if m_badsem:
        add("Marks", "warn", f"{m_badsem} row(s) have no valid semester number; skipped.")
    if bad_grade:
        add("Marks", "warn", f"{bad_grade} grade(s) not recognised; treated as blank.")
    add("Attendance", "ok", f"{len(A):,} weekly records. This semester: up to week {w_now}. "
        f"Last semester: {w_last} weeks. The semester before: {w_prev} weeks.")
    if a_unknown:
        add("Attendance", "warn", f"{a_unknown} row(s) have a roll number that is not in Students; skipped.")
    if a_badsem + a_badnum:
        add("Attendance", "warn", f"{a_badsem + a_badnum} row(s) have a missing semester, week or class count; skipped.")
    if problems:
        return result

    W = w_now
    wps = max(w_last, w_prev, W + 1)
    files = {}
    name = S.set_index("roll")["Name"]
    email = S.set_index("roll")["Email"].astype(str).str.strip().str.lower()
    email = email.where(email.str.contains("@"), "")
    files["01_sis_student_master.csv"] = pd.DataFrame({
        "Roll No": S["roll"], "Student Name": S["Name"], "Gender": S["Gender"], "Dept": S["dept"], "Batch": np.nan,
        "Year": S["year"], "Sec": S["Section"], "Residence": S["Hostel or day scholar"],
        "First Graduate": S["First graduate"].map(yes_no), "Mentor": S["Mentor"].fillna("Not assigned"),
        "HSC %": num(S["12th mark %"]).values, "Official Email": S["roll"].map(email).values,
        "Current Semester": S["cur"].astype(int)})
    files["02_exam_cell_results.csv"] = pd.DataFrame({
        "register_no": M["roll"].values, "student_name": M["roll"].map(name).values,
        "sem_label": ("Sem " + M["sem"].astype(int).astype(str)).values, "semester_no": M["sem"].astype(int).values,
        "subject_code": M["Subject code"].astype(str).str.strip().values, "subject_name": M["Subject name"].values,
        "cia1": cia1.values, "cia2": (num(M["Internal test 2"]) * scale).round(1).values,
        "end_sem": num(M["End semester"]).values, "grade": M["grade"].values, "attempt": M["attempt"].values})
    files["03_attendance_weekly.csv"] = pd.DataFrame({
        "RollNo": A["roll"].values, "semester": A["semester"].values, "week": A["week"].astype(int).values,
        "subject_code": A["Subject code"].fillna("ALL").astype(str).values,
        "classes_held": A["held"].values, "classes_attended": A["attd"].values})

    # ---- optional sheets
    def only_known(sheet):
        d = std[sheet]
        unknown = int((~d["roll"].isin(set(S["roll"]))).sum())
        if unknown:
            add(sheet, "warn", f"{unknown} row(s) have a roll number that is not in Students; skipped.")
        return d[d["roll"].isin(set(S["roll"]))].copy()

    if "LMS" in std:
        L = only_known("LMS")
        L["rel"] = num(L["Semester"]).values - L["roll"].map(cur).values
        L["week"] = num(L["Week"]).values
        L = L[L["rel"].isin([-2, -1, 0]) & L["week"].notna()]
        if len(L):
            files["04a_lms_users.csv"] = pd.DataFrame({
                "lms_user_id": S["roll"], "login_email": S["roll"].map(email).values, "display_name": S["Name"],
                "enrolled_courses": ""})
            files["04b_lms_weekly_activity.csv"] = pd.DataFrame({
                "lms_user_id": L["roll"].values, "semester": L["rel"].map(LABELS).values,
                "week": L["week"].astype(int).values, "logins": num(L["Logins"]).values,
                "minutes": num(L["Minutes online"]).values, "assignments_due": num(L["Assignments due"]).values,
                "assignments_submitted": num(L["Assignments submitted"]).values,
                "quiz_attempts": num(L["Quizzes taken"]).values})
            add("LMS", "ok", f"{len(L):,} weekly records for {L['roll'].nunique():,} students.")
        else:
            add("LMS", "skip", "No rows for the last three semesters, so it was skipped.")

    if "Other details" in std:
        O = only_known("Other details").drop_duplicates("roll").set_index("roll")
        got = []
        apt, cod = num(O["Aptitude score"]).values, num(O["Coding score"]).values
        mocks, rating = num(O["Mock interviews done"]).values, num(O["Mock interview rating"]).values
        rows = []
        for r, a, c, m, rt in zip(O.index, apt, cod, mocks, rating):
            if not np.isnan(a):
                rows.append((r, "Aptitude", "", a))
            if not np.isnan(c):
                rows.append((r, "Coding", "", c))
            for i in range(int(m) if not np.isnan(m) else 0):  # distinct rows, so cleaning keeps every interview
                rows.append((r, "Mock Interview", f"interview {i + 1}", rt))
        if rows:
            files["06_placement_training_scores.csv"] = pd.DataFrame(
                rows, columns=["candidate_id", "test_type", "attempt_month", "score"])
            got.append("placement training")
        tech, soft = num(O["Technical skill score"]), num(O["Soft skill score"])
        if tech.notna().any() or soft.notna().any():
            files["07_skills_assessment.csv"] = pd.DataFrame({
                "student_ref": O.index, "technical_score": tech.values, "soft_skill_score": soft.values,
                "assessed_on": ""})
            got.append("skills")
        ev = []
        for col, typ, label in [("Events and clubs", "Club/Event", "Event or club"),
                                ("Hackathons", "Hackathon", "Hackathon"),
                                ("Certifications", "Certification", "Certification")]:
            for r, n in zip(O.index, num(O[col]).values):
                ev += [(r, typ, label)] * (int(n) if not np.isnan(n) else 0)
        if any(num(O[c]).notna().any() for c in ("Events and clubs", "Hackathons", "Certifications")):
            E = pd.DataFrame(ev, columns=["roll_no", "event_type", "event_name"])
            files["05_student_affairs_events.csv"] = pd.DataFrame({
                "event_month": "Last 12 months", "event_type": E["event_type"], "event_name": E["event_name"],
                "participant_name": E["roll_no"].map(name).values, "department": "", "year_of_study": "",
                "roll_no": E["roll_no"]})
            got.append("activities")
        fr = num(O["Faculty rating"])
        if fr.notna().any():
            keep = fr.notna().values
            files["08a_faculty_feedback.csv"] = pd.DataFrame({
                "student": O.index[keep], "subject_code": "ALL", "participation_rating": fr.values[keep],
                "remark": O["Faculty remark"].values[keep], "given_in_week": W})
            got.append("faculty feedback")
        cs, fs, ws = num(O["Course satisfaction"]), num(O["Feels supported"]), num(O["Workload stress"])
        keep = (cs.notna() | fs.notna()).values
        if keep.any():
            files["08b_student_survey.csv"] = pd.DataFrame({
                "respondent_email": O.index[keep], "course_satisfaction": cs.fillna(fs).values[keep],
                "feels_supported": fs.fillna(cs).values[keep], "workload_stress": ws.values[keep],
                "submitted_week": W})
            got.append("survey")
        if got:
            add("Other details", "ok", f"{len(O):,} students. Found: {', '.join(got)}.")
        else:
            add("Other details", "skip", "All columns are blank, so it was skipped.")

    if "Past actions" in std:
        P = only_known("Past actions")
        P["rel"] = num(P["Semester"]).values - P["roll"].map(cur).values
        P = P[P["rel"].isin([-2, -1])]
        if len(P):
            files["09_intervention_log_past.csv"] = pd.DataFrame({
                "roll_no": P["roll"].values, "semester": P["rel"].map(LABELS).values, "week": num(P["Week"]).values,
                "action": P["Action"].astype(str).str.strip().values, "owner": P["Done by"].values})
            add("Past actions", "ok", f"{len(P):,} actions from the last two semesters.")
        else:
            add("Past actions", "skip", "No actions from the last two semesters, so it was skipped.")

    if "Past placements" in std:
        G = std["Past placements"].copy()
        G["placed"] = G["Placed"].map(yes_no)
        G["cgpa"] = num(G["Final CGPA"]).values
        G = G[G["placed"].notna() & G["cgpa"].notna()]
        vc = G["placed"].value_counts()
        if len(G) >= 20 and vc.get("Yes", 0) >= 5 and vc.get("No", 0) >= 5:
            files["10_alumni_placement_outcomes.csv"] = pd.DataFrame({
                "alumni_id": G["Student ID"].values, "dept": G["Department"].map(clean_dept).values,
                "final_cgpa": G["cgpa"].values, "standing_backlogs": num(G["Standing arrears"]).fillna(0).values,
                "best_aptitude": num(G["Aptitude score"]).values, "best_coding": num(G["Coding score"]).values,
                "technical_score": num(G["Technical skill score"]).values,
                "soft_skill_score": num(G["Soft skill score"]).values,
                "mock_interviews": num(G["Mock interviews done"]).values,
                "avg_mock_rating": num(G["Mock interview rating"]).values,
                "certifications": num(G["Certifications"]).values,
                "internship": G["Internship"].map(yes_no).map({"Yes": 1, "No": 0}).values,
                "placed": G["placed"].values, "offer_tier": ""})
            add("Past placements", "ok", f"{len(G):,} graduates, {vc.get('Yes', 0)} placed.")
        else:
            add("Past placements", "warn", "Needs at least 20 graduates, with at least 5 placed and 5 not placed. "
                                           "Using the built-in sample history instead.")

    depts = {d: [d, f"U{i:02d}", f"D{i}", int(n)] for i, (d, n) in enumerate(S["dept"].value_counts().items())}
    result.update(ok=True, files=files, settings=dict(
        semesters=[LABELS[-2], LABELS[-1], LABELS[0]], current_week=int(W), weeks_per_sem=int(wps), depts=depts),
        facts=dict(students=len(S), depts=n_depts, week=int(W), weeks_per_sem=int(wps)))
    return result


# ---------------------------------------------------------------- the template workbook
def _styled_workbook(sheets_rows):
    """sheets_rows: {sheet: list of rows (lists)}. Returns xlsx bytes with a 'Start here' sheet first."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    crimson, grey, ink = "B3163A", "E9ECF2", "1D2433"
    wb = Workbook()
    ws = wb.active
    ws.title = "Start here"
    ws.sheet_properties.tabColor = ink
    lines = [
        ("EarlySignal: data template", Font(bold=True, size=16, color=crimson)),
        ("", None),
        ("1. Fill one sheet for each kind of record. Only the 3 red sheets are needed: Students, Marks, Attendance.", None),
        ("2. Use the same Roll No in every sheet.", None),
        ("3. Semester means the semester number, 1 to 8. Include this semester and the two before it.", None),
        ("4. Keep the column headings. Columns marked * are needed; any other column can be left blank.", None),
        ("5. Rows whose Roll No starts with EXAMPLE are ignored, so you can paste below them or delete them.", None),
        ("6. Then open EarlySignal > Data sources > Excel upload and upload this file. Separate CSVs per sheet also work.", None),
        ("", None),
    ]
    for text, font in lines:
        ws.append([text])
        if font:
            ws.cell(ws.max_row, 1).font = font
    ws.append(["Sheet", "Needed?", "Rows", "Used for"])
    for c in range(1, 5):
        ws.cell(ws.max_row, c).font = Font(bold=True, color="FFFFFF")
        ws.cell(ws.max_row, c).fill = PatternFill("solid", fgColor=ink)
    for s, v in SHEETS.items():
        ws.append([s, "Needed" if v["required"] else "Optional", v["per"], v["why"]])
        if v["required"]:
            ws.cell(ws.max_row, 2).font = Font(bold=True, color=crimson)
    ws.append([])
    ws.append(["Sheet", "Column", "Needed?", "What to put"])
    for c in range(1, 5):
        ws.cell(ws.max_row, c).font = Font(bold=True, color="FFFFFF")
        ws.cell(ws.max_row, c).fill = PatternFill("solid", fgColor=ink)
    for s, v in SHEETS.items():
        for col, (need, help_, _) in v["cols"].items():
            ws.append([s, col, "Needed" if need else "", help_])
    for col, w in zip("ABCD", (24, 26, 12, 90)):
        ws.column_dimensions[col].width = w

    for s, v in SHEETS.items():
        sh = wb.create_sheet(s)
        sh.sheet_properties.tabColor = crimson if v["required"] else "9AA3B2"
        heads = [c + (" *" if need else "") for c, (need, _, _) in v["cols"].items()]
        sh.append(heads)
        for i, (c, (need, _, _)) in enumerate(v["cols"].items(), start=1):
            cell = sh.cell(1, i)
            cell.font = Font(bold=True, color="FFFFFF" if need else ink)
            cell.fill = PatternFill("solid", fgColor=crimson if need else grey)
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            sh.column_dimensions[get_column_letter(i)].width = max(12, min(26, len(c) + 4))
        sh.freeze_panes = "A2"
        for row in sheets_rows.get(s, []):
            sh.append([None if (isinstance(x, float) and np.isnan(x)) else x for x in row])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def template_bytes():
    return _styled_workbook(EXAMPLES)


def example_bytes(dept="CSE"):
    """A filled workbook built from the sample college (one department), to show what real data looks like."""
    from . import config as C
    from .integrate import integrate

    with C.sample_mode():
        d = integrate()
    sis = d["sis"][d["sis"]["dept"] == dept].copy()
    sis["cur"] = 2 * sis["year"] - 1  # the sample college is in an odd semester
    rolls = set(sis["roll_no"])
    cur = sis.set_index("roll_no")["cur"]
    rel = {s: k - 2 for k, s in enumerate(["2025-ODD", "2026-EVEN", "2026-ODD"])}
    rows = {}
    rows["Students"] = sis[["roll_no", "name", "dept", "cur", "section", "gender", "residence", "first_graduate",
                            "mentor", "email", "hsc_percent"]].values.tolist()
    ex = d["exams"][d["exams"]["roll_no"].isin(rolls)].copy()
    ex["cia1"] = np.where(ex["cia1_absent"], "AB", ex["cia1"])
    rows["Marks"] = ex[["roll_no", "semester_no", "subject_code", "subject_name", "cia1", "cia2", "end_sem", "grade",
                        "attempt"]].values.tolist()
    at = d["attendance"][d["attendance"]["roll_no"].isin(rolls)].copy()
    at["sem"] = at["roll_no"].map(cur) + at["semester"].map(rel)
    rows["Attendance"] = at[["roll_no", "sem", "week", "subject_code", "classes_held",
                             "classes_attended"]].values.tolist()
    lm = d["lms"][d["lms"]["roll_no"].isin(rolls)].copy()
    lm["sem"] = lm["roll_no"].map(cur) + lm["semester"].map(rel)
    rows["LMS"] = lm[["roll_no", "sem", "week", "logins", "minutes", "assignments_due", "assignments_submitted",
                      "quiz_attempts"]].values.tolist()
    pl = d["placement"]
    best = pl.groupby(["roll_no", "test_type"])["score_num"].max().unstack()
    mock = pl[pl["test_type"] == "Mock Interview"].groupby("roll_no")["score_num"].agg(["size", "mean"])
    ev = d["events"]
    cnt = lambda t: ev[ev["event_type"] == t].groupby("roll_no").size()  # noqa: E731
    events = ev[~ev["event_type"].isin(["Certification", "Hackathon"])].groupby("roll_no").size()
    sk = d["skills"].set_index("roll_no")
    fb = d["feedback"].groupby("roll_no")
    sv = d["survey"].set_index("roll_no")
    O = pd.DataFrame(index=sorted(rolls))
    O["apt"], O["cod"] = best.get("Aptitude"), best.get("Coding")
    O["mocks"], O["rating"] = mock["size"], mock["mean"].round(1)
    O["tech"], O["soft"] = sk["technical_score"], sk["soft_skill_score"]
    O["cert"], O["events"], O["hack"] = cnt("Certification"), events, cnt("Hackathon")
    O[["cert", "events", "hack"]] = O[["cert", "events", "hack"]].fillna(0)
    O["frate"] = fb["participation_rating"].mean().round(1)
    O["remark"] = d["feedback"].sort_values("participation_rating").groupby("roll_no")["remark"].first()
    O["sat"], O["sup"], O["stress"] = sv["course_satisfaction"], sv["feels_supported"], sv["workload_stress"]
    O.loc[O.index.isin(sis.loc[sis["year"] == 2, "roll_no"]), "mocks"] = O["mocks"].fillna(0)
    rows["Other details"] = O.reset_index().values.tolist()
    iv = d["interventions"][d["interventions"]["roll_no"].isin(rolls)].copy()
    iv["sem"] = iv["roll_no"].map(cur) + iv["semester"].map(rel)
    rows["Past actions"] = iv[["roll_no", "sem", "week", "action", "owner"]].values.tolist()
    al = d["alumni"]
    rows["Past placements"] = al[["alumni_id", "dept", "final_cgpa", "standing_backlogs", "best_aptitude",
                                  "best_coding", "technical_score", "soft_skill_score", "mock_interviews",
                                  "avg_mock_rating", "certifications", "internship", "placed"]].values.tolist()
    for r in rows["Past placements"]:
        r[11] = "Yes" if r[11] in (1, "1", "Yes", True) else "No"
    return _styled_workbook(rows)
